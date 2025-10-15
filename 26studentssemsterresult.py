from docx import Document
class Student:
    def __init__(self, reg_no, name, dept, year):
        self.reg_no = reg_no
        self.name = name
        self.dept = dept
        self.year = year

    def save_personal_info(self):
        filename = f"studentreg-{self.reg_no}.txt"
        with open(filename, "w") as f:
            f.write(f"Register No: {self.reg_no}\n")
            f.write(f"Name: {self.name}\n")
            f.write(f"Department: {self.dept}\n")
            f.write(f" Personal info saved to {filename}")
class Assessment:
    def __init__(self, reg_no, subject_marks):
        self.reg_no = reg_no
        self.subject_marks = subject_marks  

    def save_assessment(self):
        filename = f"studentreg-{self.reg_no}-marks.txt"
        with open(filename, "w") as f:
            total = 0
            for subject, mark in self.subject_marks.items():
                f.write(f"{subject}: {mark}\n")
                total += mark
            avg = total / len(self.subject_marks)
            f.write(f"\nTotal: {total}\nAverage: {avg:.2f}\n")
        print(f"Assessment saved to {filename}")
class Report:
    def __init__(self, students):
        self.students = students  

    def generate_word_report(self):
        doc = Document()
        doc.add_heading("STUDENT PERFORMANCE REPORT", level=1)

        for student, assessment in self.students:
            total = sum(assessment.subject_marks.values())
            avg = total / len(assessment.subject_marks)

            doc.add_heading(f"{student.name} ({student.reg_no})", level=2)
            doc.add_paragraph(f"Department: {student.dept}")
            doc.add_paragraph(f"Year: {student.year}")

            table = doc.add_table(rows=1, cols=2)
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = "Subject"
            hdr_cells[1].text = "Marks"

            for subject, mark in assessment.subject_marks.items():
                row_cells = table.add_row().cells
                row_cells[0].text = subject
                row_cells[1].text = str(mark)

            doc.add_paragraph(f"Total: {total}")
            doc.add_paragraph(f"Average: {avg:.2f}")
            doc.add_paragraph("---------------------------------------")

        filename = "student_analysis_report.docx"
        doc.save(filename)
        print(f" Word report saved as '{filename}'")
if __name__ == "__main__":
    n = int(input("Enter number of students: "))
    all_students = []

    for i in range(n):
        print(f"\nEnter details for Student {i+1}")
        reg_no = input("Register No: ")
        name = input("Name: ")
        dept = input("Department: ")
        year = input("Year: ")

        student = Student(reg_no, name, dept, year)
        student.save_personal_info()

        print("\nEnter marks for 3 subjects:")
        subject_marks = {}
        for sub in ["Maths", "Science", "English"]:
            mark = int(input(f"{sub}: "))
            subject_marks[sub] = mark

        assessment = Assessment(reg_no, subject_marks)
        assessment.save_assessment()

        all_students.append((student, assessment))

    report = Report(all_students)
    report.generate_word_report()
