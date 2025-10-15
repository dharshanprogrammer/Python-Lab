import qrcode

from docx import Document


class Vehicle:



    def __init__(self, reg_no, owner_name, model, color, engine_no, chassis_no):
        self.reg_no = reg_no
        self.owner_name = owner_name
        self.model = model
        self.color = color
        self.engine_no = engine_no
        self.chassis_no = chassis_no

    def save_vehicle_details(self):

        doc = Document()
        doc.add_heading('Vehicle Registration Details', 0)

        doc.add_paragraph(f"Registration Number: {self.reg_no}")
        doc.add_paragraph(f"Owner Name: {self.owner_name}")
        doc.add_paragraph(f"Vehicle Model: {self.model}")
        doc.add_paragraph(f"Color: {self.color}")
        doc.add_paragraph(f"Engine Number: {self.engine_no}")
        doc.add_paragraph(f"Chassis Number: {self.chassis_no}")


        filename = f"vehicle_{self.reg_no}.docx"
        doc.save(filename)
        print(f" Vehicle details successfully saved in a true DOCX file: '{filename}'")

    def generate_rc_book(self):

        qr_data = (
            f"Reg No: {self.reg_no}\n"
            f"Owner: {self.owner_name}\n"
            f"Model: {self.model}\n"
            f"Color: {self.color}\n"
            f"Engine No: {self.engine_no}\n"
            f"Chassis No: {self.chassis_no}"
        )

        qr = qrcode.QRCode(version=1, box_size=10, border=4)
        qr.add_data(qr_data)
        qr.make(fit=True)
        qr_filename = f"qr_{self.reg_no}.png"
        qr_img = qr.make_image(fill_color="black", back_color="white")
        qr_img.save(qr_filename)

        # --- HTML Generation (Digital RC Book) ---
        html_content = f"""
        <html>
        <head>
            <title>Digital RC Book - {self.reg_no}</title>
            <style>
                body {{
                    font-family: Arial, sans-serif;
                    background-color: #f4f4f4;
                    padding: 20px;
                }}
                .rcbook {{
                    background: white;
                    padding: 20px;
                    border-radius: 10px;
                    box-shadow: 0px 0px 10px rgba(0,0,0,0.2);
                    width: 500px;
                    margin: auto;
                }}
                img {{
                    display: block;
                    margin: 20px auto;
                }}
                h1 {{
                    text-align: center;
                    color: #333;
                }}
                table {{
                    width: 100%;
                    border-collapse: collapse;
                }}
                td {{
                    padding: 8px;
                }}
            </style>
        </head>
        <body>
            <div class="rcbook">
                <h1>Digital RC Book</h1>
                <table>
                    <tr><td><b>Registration Number:</b></td><td>{self.reg_no}</td></tr>
                    <tr><td><b>Owner Name:</b></td><td>{self.owner_name}</td></tr>
                    <tr><td><b>Vehicle Model:</b></td><td>{self.model}</td></tr>
                    <tr><td><b>Color:</b></td><td>{self.color}</td></tr>
                    <tr><td><b>Engine Number:</b></td><td>{self.engine_no}</td></tr>
                    <tr><td><b>Chassis Number:</b></td><td>{self.chassis_no}</td></tr>
                </table>
                <h3 style="text-align:center;">Scan QR Code for Verification</h3>
                <img src="{qr_filename}" width="150" height="150" />
            </div>
        </body>
        </html>
        """

        html_filename = f"RC_{self.reg_no}.html"
        with open(html_filename, "w", encoding="utf-8") as file:
            file.write(html_content)

        print(f"📄 Digital RC Book created: '{html_filename}'")
        print(f"🔗 QR Code saved as: '{qr_filename}'")


if __name__ == "__main__":
    print("=== VEHICLE REGISTRATION SYSTEM ===")
    reg_no = input("Enter Vehicle Registration Number: ")
    owner_name = input("Enter Owner Name: ")
    model = input("Enter Vehicle Model: ")
    color = input("Enter Color: ")
    engine_no = input("Enter Engine Number: ")
    chassis_no = input("Enter Chassis Number: ")

    v = Vehicle(reg_no, owner_name, model, color, engine_no, chassis_no)
    v.save_vehicle_details()
    v.generate_rc_book()
